"""Genera 'documentos/Tabla de precios.xlsx' para que los dueños revisen y aprueben los precios.

Hojas:
  1. Tabla de precios  -> para los dueños: rangos, ejemplos, costo, venta, reparto y notas.
  2. Desglose          -> cálculo detallado por rango y ganancia por tóner.
  3. Supuestos         -> tóner, papel, mantenimiento, energía, factor, % ahorro (celdas editables).

Todo el libro usa fórmulas: si se cambia un supuesto en Excel, se recalcula solo.
Los valores salen de la misma configuración que usa el programa (config.json si existe).

Uso:  .venv\\Scripts\\python herramientas\\generar_tabla_precios.py
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
from cotizador import config as _cfg  # noqa: E402
from cotizador.costos import tabla_rangos  # noqa: E402

CFG = _cfg.cargar()  # misma configuración que el programa

SALIDA = RAIZ / "documentos" / "Tabla de precios.xlsx"

# Cobertura típica de cada rango (fracción de la hoja). Medida con los PDFs de
# muestras/referencia (herramientas/generar_referencias.py), promedio de 4 fotos.
# Los rangos b/n cargado/total no tienen referencia: se usa el punto medio del rango.
TIPICO = {
    "B/N normal":   (0.000, 0.070),  # (C+M+Y, K)  hoja de texto completa
    "B/N cargado":  (0.000, 0.260),
    "B/N total":    (0.000, 0.700),
    "Color mínimo": (0.057, 0.076),  # texto + título decorado
    "Color medio":  (0.406, 0.155),  # foto de media página
    "Color alto":   (0.813, 0.310),  # dos fotos
    "Color total":  (1.158, 0.443),  # foto a página completa
}
EJEMPLOS = {
    "B/N normal":   "Texto, tareas, cartas, formularios",
    "B/N cargado":  "Tablas oscuras, gráficas o fotos en gris",
    "B/N total":    "Foto en gris a página completa, fondos negros",
    "Color mínimo": "Documento con título, logo o firma a color",
    "Color medio":  "Una foto o imagen de media página",
    "Color alto":   "Dos fotos o una imagen grande",
    "Color total":  "Foto o afiche a página completa",
}
RELLENOS = {
    "B/N normal": "F2F2F2", "B/N cargado": "D9D9D9", "B/N total": "BFBFBF",
    "Color mínimo": "FFF4D6", "Color medio": "FFE0B2", "Color alto": "FFCC99", "Color total": "F8B195",
}

FUENTE = "Arial"
AZUL, VERDE = "0000FF", "008000"
PESOS = '"$"#,##0'
PCT = "0.0%"
DELGADO = Side(style="thin", color="A6A6A6")
BORDE = Border(left=DELGADO, right=DELGADO, top=DELGADO, bottom=DELGADO)
AMARILLO = PatternFill("solid", fgColor="FFFF00")
CABECERA = PatternFill("solid", fgColor="1F4E78")


def f(color="000000", bold=False, size=10, italic=False):
    return Font(name=FUENTE, color=color, bold=bold, size=size, italic=italic)


def cabecera(ws, fila, textos, col_ini=1):
    for i, t in enumerate(textos):
        c = ws.cell(fila, col_ini + i, t)
        c.font = f("FFFFFF", bold=True)
        c.fill = CABECERA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE


# ============================================================================ Supuestos
def hoja_supuestos(wb):
    ws = wb.create_sheet("Supuestos")
    ws["A1"] = "Supuestos del cálculo"
    ws["A1"].font = f(bold=True, size=14)
    ws["A2"] = "Celdas en azul con fondo amarillo = datos que se pueden cambiar. Todo lo demás se recalcula solo."
    ws["A2"].font = f(italic=True)

    t, p = CFG["toner"], CFG["papel"]
    filas = [
        ("Tóner", None, None, None),
        ("Precio de un tóner de color (C, M o Y)", t["C"]["precio"], PESOS, "Dato de Felipe, sep. 2026"),
        ("Rendimiento tóner de color (hojas al 5 %)", t["C"]["rendimiento"], "#,##0", "Estimado de Felipe"),
        ("Precio del tóner negro (K)", t["K"]["precio"], PESOS, "Dato de Felipe, sep. 2026"),
        ("Rendimiento tóner negro (hojas al 5 %)", t["K"]["rendimiento"], "#,##0", "Estimado de Felipe"),
        ("Cobertura con que se mide el rendimiento", CFG["cobertura_referencia_pct"] / 100, PCT, "Norma de los fabricantes"),
        ("Factor de corrección (seguridad en tinta)", CFG["factor_correccion"], "0.00", "1,15 = se asume 15 % más tinta de la calculada"),
        ("Papel", None, None, None),
        ("Resma carta", p["carta"]["precio_resma"], PESOS, "Dato de Felipe"),
        ("Resma oficio", p["oficio"]["precio_resma"], PESOS, "Dato de Felipe"),
        ("Hojas por resma", p["carta"]["hojas_resma"], "#,##0", "Dato de Felipe"),
        ("Alto hoja carta (mm)", p["carta"]["alto_mm"], "0.0", "Mismo ancho que oficio (215,9 mm)"),
        ("Alto hoja oficio (mm)", p["oficio"]["alto_mm"], "0.0", "Oficio 21,6 × 33 cm"),
        ("Energía y volumen", None, None, None),
        ("Potencia promedio encendida (W)", CFG["energia"]["potencia_w"], "#,##0", "Alto a propósito: la ficha Ricoh da 1,16 kWh/semana"),
        ("Horas encendida al día", CFG["energia"]["horas_dia"], "0.0", ""),
        ("Días de trabajo al mes", CFG["energia"]["dias_mes"], "0", ""),
        ("Precio del kWh", CFG["energia"]["precio_kwh"], PESOS, "EPM estrato 4: $885 (nov. 2026) + margen por recargos"),
        ("Hojas impresas al mes", CFG["hojas_mes"], "#,##0", "ESTIMADO: ver contador de la Ricoh"),
        ("Reparto de lo que sobra", None, None, None),
        ("Parte para ahorro (imprevistos, reposición)", CFG["capital_pct"] / 100, PCT, "El resto es ganancia"),
    ]
    fila = 4
    celdas = {}
    for etiqueta, valor, fmt, nota in filas:
        if valor is None:
            ws.cell(fila, 1, etiqueta).font = f(bold=True, size=11)
        else:
            ws.cell(fila, 1, etiqueta).font = f()
            c = ws.cell(fila, 2, valor)
            c.font, c.number_format, c.fill, c.border = f(AZUL), fmt, AMARILLO, BORDE
            ws.cell(fila, 3, nota).font = f(italic=True, color="595959")
            celdas[etiqueta] = f"Supuestos!$B${fila}"
        fila += 1

    # ---- Mantenimiento: tabla con nombre, costo, ciclo y si solo aplica a color
    fila += 1
    ws.cell(fila, 1, "Mantenimiento (cada repuesto o visita se reparte entre las hojas de su ciclo)").font = f(bold=True, size=11)
    fila += 1
    cabecera(ws, fila, ["Repuesto o visita", "Costo", "Cada (hojas)", "¿Solo color?", "Por hoja"])
    fila += 1
    ini_mant = fila
    for m in CFG["mantenimiento"]:
        ws.cell(fila, 1, m["nombre"]).font = f(AZUL)
        for col, v, fmt in ((2, m["costo"], PESOS), (3, m["cada_hojas"], "#,##0"), (4, "Sí" if m["solo_color"] else "No", None)):
            c = ws.cell(fila, col, v)
            c.font, c.fill, c.border = f(AZUL), AMARILLO, BORDE
            if fmt:
                c.number_format = fmt
        c = ws.cell(fila, 5, f"=B{fila}/C{fila}")
        c.number_format, c.border, c.font = '"$"#,##0.0', BORDE, f()
        fila += 1
    fin_mant = fila - 1
    ws.cell(fila, 1, "ESTIMADOS: confirmar costos y ciclos con el técnico. Ciclo de 120.000 hojas: típico de kits de fusor Ricoh."
            ).font = f(italic=True, color="595959")
    rango_b, rango_c, rango_d = (f"Supuestos!${col}${ini_mant}:${col}${fin_mant}" for col in "BCD")

    fila += 2
    ws.cell(fila, 1, "Valores calculados").font = f(bold=True, size=11)
    fila += 1
    calc = [
        ("tinta_color", "Costo de tinta de color por hoja carta al 100 %",
         f"={celdas['Precio de un tóner de color (C, M o Y)']}/{celdas['Rendimiento tóner de color (hojas al 5 %)']}"
         f"/{celdas['Cobertura con que se mide el rendimiento']}", "precio ÷ rendimiento ÷ 5 %"),
        ("tinta_negra", "Costo de tinta negra por hoja carta al 100 %",
         f"={celdas['Precio del tóner negro (K)']}/{celdas['Rendimiento tóner negro (hojas al 5 %)']}"
         f"/{celdas['Cobertura con que se mide el rendimiento']}", "precio ÷ rendimiento ÷ 5 %"),
        ("papel_carta", "Papel por hoja carta", f"={celdas['Resma carta']}/{celdas['Hojas por resma']}", ""),
        ("papel_oficio", "Papel por hoja oficio", f"={celdas['Resma oficio']}/{celdas['Hojas por resma']}", ""),
        ("area_oficio", "Área oficio ÷ área carta",
         f"={celdas['Alto hoja oficio (mm)']}/{celdas['Alto hoja carta (mm)']}", "la tinta crece con el área"),
        ("mant_bn", "Mantenimiento por hoja b/n", f'=SUMPRODUCT({rango_b}/{rango_c}*({rango_d}="No"))',
         "suma de repuestos que no son solo de color"),
        ("mant_color", "Mantenimiento por hoja a color", f"=SUMPRODUCT({rango_b}/{rango_c})", "todos los repuestos"),
        ("kwh_mes", "Energía al mes (kWh)",
         f"={celdas['Potencia promedio encendida (W)']}/1000*{celdas['Horas encendida al día']}*{celdas['Días de trabajo al mes']}", ""),
        ("energia_mes", "Energía al mes ($)", f"=B{{kwh}}*{celdas['Precio del kWh']}", ""),
        ("energia_hoja", "Energía por hoja", f"=B{{emes}}/{celdas['Hojas impresas al mes']}", "la del mes ÷ hojas del mes"),
    ]
    filas_calc = {}
    for clave, etiqueta, formula, nota in calc:
        ws.cell(fila, 1, etiqueta).font = f()
        formula = formula.replace("{kwh}", str(filas_calc.get("kwh_mes", ""))).replace("{emes}", str(filas_calc.get("energia_mes", "")))
        filas_calc[clave] = fila
        c = ws.cell(fila, 2, formula)
        c.font, c.border = f(), BORDE
        c.number_format = {"area_oficio": "0.00", "kwh_mes": "#,##0"}.get(clave, PESOS)
        ws.cell(fila, 3, nota).font = f(italic=True, color="595959")
        celdas[clave] = f"Supuestos!$B${fila}"
        fila += 1

    # Alias cortos para las fórmulas de las otras hojas
    celdas["factor"] = celdas["Factor de corrección (seguridad en tinta)"]
    celdas["capital"] = celdas["Parte para ahorro (imprevistos, reposición)"]
    celdas["toner_color"] = celdas["Precio de un tóner de color (C, M o Y)"]
    celdas["toner_negro"] = celdas["Precio del tóner negro (K)"]

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 46
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 11
    return celdas


# ============================================================================ Desglose
def hoja_desglose(wb, s):
    ws = wb.create_sheet("Desglose")
    ws["A1"] = "Desglose de costos por rango (hoja carta)"
    ws["A1"].font = f(bold=True, size=14)
    ws["A2"] = ("Azul = dato editable. Negro = fórmula. Verde = viene de la hoja Supuestos. "
                "Cobertura típica medida con páginas de referencia; límite = máximo del rango en el programa.")
    ws["A2"].font = f(italic=True)

    cols = ["Rango", "Cobertura típica color (C+M+Y)", "Cobertura típica negro (K)", "Límite de cobertura del rango",
            "Tinta por hoja (típica)", "Papel por hoja", "Mantenimiento + energía", "Costo aproximado (todo)",
            "Costo máximo del rango", "Precio de venta", "Sobrante (venta − costo)", "Para ahorro", "Ganancia",
            "Costo aproximado en oficio", "Ganancia en oficio"]
    cabecera(ws, 4, cols)
    ws.row_dimensions[4].height = 48

    rangos = CFG["rangos_bn"] + CFG["rangos_color"]
    filas = {}
    for i, r in enumerate(rangos):
        n = 5 + i
        nombre = r["nombre"]
        es_color = nombre.startswith("Color")
        filas[nombre] = n
        cmy, k = TIPICO[nombre]
        tinta_max = s["tinta_color"] if es_color else s["tinta_negra"]
        mant = s["mant_color"] if es_color else s["mant_bn"]
        precio = tabla_rangos(CFG, "carta", es_color)[[x["nombre"] for x in CFG["rangos_color" if es_color else "rangos_bn"]].index(nombre)]["precio"]
        valores = [
            (nombre, None, f(bold=True)),
            (cmy, PCT, f(AZUL)),
            (k, PCT, f(AZUL)),
            (r["cobertura_max_pct"] / 100, PCT, f(AZUL)),
            (f"=(B{n}*{s['tinta_color']}+C{n}*{s['tinta_negra']})*{s['factor']}", PESOS, f()),
            (f"={s['papel_carta']}", PESOS, f(VERDE)),
            (f"={mant}+{s['energia_hoja']}", PESOS, f(VERDE)),
            (f"=E{n}+F{n}+G{n}", PESOS, f(bold=True)),
            (f"=D{n}*{tinta_max}*{s['factor']}+F{n}+G{n}", PESOS, f()),
            (precio, PESOS, f(AZUL, bold=True)),
            (f"=J{n}-H{n}", PESOS, f()),
            (f"=K{n}*{s['capital']}", PESOS, f()),
            (f"=K{n}-L{n}", PESOS, f(bold=True)),
            (f"=E{n}*{s['area_oficio']}+{s['papel_oficio']}+G{n}", PESOS, f()),
            (f"=(J{n}-N{n})*(1-{s['capital']})", PESOS, f()),
        ]
        for j, (v, fmt, fuente) in enumerate(valores, start=1):
            c = ws.cell(n, j, v)
            c.font, c.border = fuente, BORDE
            if fmt:
                c.number_format = fmt
            c.alignment = Alignment(horizontal="left" if j == 1 else "center")
        ws.cell(n, 1).fill = PatternFill("solid", fgColor=RELLENOS[nombre])
    ultima = 5 + len(rangos) - 1

    # ---- Ganancia por tóner
    ini = ultima + 3
    ws.cell(ini, 1, "¿Cuánto deja cada tóner?").font = f(bold=True, size=12)
    ws.cell(ini + 1, 1, ("Si todo el tóner se gastara en hojas de un solo rango: cuántas hojas alcanza a pagar "
                         "la tinta de un tóner, y cuánto entra por ellas. Incluye el factor de corrección.")).font = f(italic=True)
    cab = ["Rango", "Precio de un tóner", "Hojas que paga ese tóner", "Ventas",
           "Papel, mantenimiento y energía", "Para ahorro", "Ganancia"]
    cabecera(ws, ini + 2, cab)
    ws.row_dimensions[ini + 2].height = 32
    for i, r in enumerate(rangos):
        n = ini + 3 + i
        d = filas[r["nombre"]]
        es_color = r["nombre"].startswith("Color")
        toner = s["toner_color"] if es_color else s["toner_negro"]
        valores = [
            (f"=A{d}", None, f(bold=True)),
            (f"={toner}", PESOS, f(VERDE)),
            (f"=B{n}/E{d}", "#,##0", f()),
            (f"=C{n}*J{d}", PESOS, f()),
            (f"=C{n}*(F{d}+G{d})", PESOS, f()),
            (f"=C{n}*L{d}", PESOS, f()),
            (f"=C{n}*M{d}", PESOS, f(bold=True)),
        ]
        for j, (v, fmt, fuente) in enumerate(valores, start=1):
            c = ws.cell(n, j, v)
            c.font, c.border = fuente, BORDE
            if fmt:
                c.number_format = fmt
            c.alignment = Alignment(horizontal="left" if j == 1 else "center")
        ws.cell(n, 1).fill = PatternFill("solid", fgColor=RELLENOS[r["nombre"]])
    fin_toner = ini + 3 + len(rangos) - 1

    nota = fin_toner + 2
    textos = [
        "Cómo leer esta tabla:",
        "• Tinta por hoja = (precio del tóner ÷ rendimiento) × (cobertura ÷ 5 %) × factor de corrección.",
        "• Mantenimiento + energía: repuestos y visitas repartidos por hoja, más la luz del mes ÷ hojas del mes (hoja Supuestos).",
        "• Sobrante = precio de venta − todos los costos. Se reparte entre ahorro y ganancia según el % de la hoja Supuestos.",
        "• Tóner: una hoja a color también gasta algo de negro; aquí se cuenta toda la tinta de la hoja contra un tóner.",
        "• No incluye arriendo ni el tiempo de trabajo.",
    ]
    for i, t in enumerate(textos):
        ws.cell(nota + i, 1, t).font = f(bold=(i == 0))

    anchos = [16, 14, 14, 14, 13, 11, 14, 14, 14, 12, 14, 12, 12, 14, 13]
    for j, w in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(4, j).column_letter].width = w
    ws.freeze_panes = "B5"
    return filas, ini + 3


# ============================================================================ Tabla de precios
def hoja_tabla(wb, filas, fila_toner):
    ws = wb.active
    ws.title = "Tabla de precios"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Tabla de precios de impresión"
    ws["A1"].font = f(bold=True, size=18, color="1F4E78")
    ws["A2"] = "Impresora Ricoh MP C3003 · precio por hoja · tamaño carta y oficio · septiembre de 2026"
    ws["A2"].font = f(size=11, color="595959")

    cab = ["Rango", "¿Qué tipo de hoja es?", "Nos cuesta\n(aprox.)", "Se cobra", "Para\nahorro", "Ganancia"]
    cabecera(ws, 4, cab)
    ws.row_dimensions[4].height = 34

    fila = 5
    for grupo, rangos in (("BLANCO Y NEGRO", CFG["rangos_bn"]), ("COLOR", CFG["rangos_color"])):
        c = ws.cell(fila, 1, grupo)
        c.font = f(bold=True, size=11, color="1F4E78")
        fila += 1
        for r in rangos:
            d = filas[r["nombre"]]
            valores = [
                (r["nombre"], None, f(bold=True, size=11)),
                (EJEMPLOS[r["nombre"]], None, f(size=10)),
                (f"=Desglose!H{d}", PESOS, f(size=11)),
                (f"=Desglose!J{d}", PESOS, f(bold=True, size=12)),
                (f"=Desglose!L{d}", PESOS, f(size=11)),
                (f"=Desglose!M{d}", PESOS, f(size=11, color="006100", bold=True)),
            ]
            for j, (v, fmt, fuente) in enumerate(valores, start=1):
                c = ws.cell(fila, j, v)
                c.font, c.border = fuente, BORDE
                c.fill = PatternFill("solid", fgColor=RELLENOS[r["nombre"]])
                c.alignment = Alignment(horizontal="left" if j <= 2 else "center", vertical="center", wrap_text=True)
                if fmt:
                    c.number_format = fmt
            ws.row_dimensions[fila].height = 22
            fila += 1
        fila += 1

    # ---- Cuánto deja un tóner (dos ejemplos)
    ws.cell(fila, 1, "¿Cuánto deja un tóner?").font = f(bold=True, size=12, color="1F4E78")
    fila += 1
    cabecera(ws, fila, ["Ejemplo", "", "Hojas", "Se vende", "Para\nahorro", "Ganancia"])
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
    ws.row_dimensions[fila].height = 30
    fila += 1
    idx = {r["nombre"]: i for i, r in enumerate(CFG["rangos_bn"] + CFG["rangos_color"])}
    ejemplos = [("Un tóner negro gastado en hojas de texto (B/N normal)", "B/N normal"),
                ("Un tóner de color gastado en hojas con una foto (Color medio)", "Color medio")]
    for texto, rango in ejemplos:
        t = fila_toner + idx[rango]
        ws.cell(fila, 1, texto).font = f(size=10)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
        for j, (col, fmt) in enumerate([("C", "#,##0"), ("D", PESOS), ("F", PESOS), ("G", PESOS)], start=3):
            c = ws.cell(fila, j, f"=Desglose!{col}{t}")
            c.number_format, c.border = fmt, BORDE
            c.font = f(size=11, bold=(col == "G"), color="006100" if col == "G" else "000000")
            c.alignment = Alignment(horizontal="center", vertical="center")
        for j in (1, 2):
            ws.cell(fila, j).border = BORDE
            ws.cell(fila, j).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[fila].height = 30
        fila += 1

    # ---- Notas
    fila += 1
    ws.cell(fila, 1, "Notas").font = f(bold=True, size=12, color="1F4E78")
    fila += 1
    notas = [
        "El programa revisa cada hoja y la pone sola en su rango según la tinta que usa. "
        "Si el trabajo tiene varias hojas, muestra cuántas hay de cada rango y el total.",
        "\"Nos cuesta\" incluye la tinta (con 15 % extra por seguridad), el papel, el mantenimiento de la "
        "impresora (técnico y repuestos) y la luz, para una hoja típica del rango. No incluye arriendo ni el "
        "tiempo de trabajo. Mantenimiento, luz y hojas al mes son estimados: ajustarlos en el programa (botón Costos).",
        "Blanco y negro se cobra $700 en los tres rangos. Los rangos se muestran para que se vea que una hoja "
        "de texto cuesta muy poco y que incluso una hoja casi toda negra sigue dejando buena ganancia.",
        "Color va de $1.000 a $4.000 según cuánta hoja ocupe el color. Una hoja de texto con un título o logo "
        "a color ya cuenta como color (rango mínimo).",
        "Lo que sobra después de pagar todos los costos se reparte mitad para ahorro (imprevistos, una impresora "
        "nueva algún día) y mitad ganancia. Se sugiere guardar el ahorro aparte.",
        "Oficio se cobra igual que carta. Una hoja oficio gasta cerca de 18 % más tinta y el papel cuesta $38 "
        "en vez de $30; como se vende poco y el margen es amplio, no vale la pena cobrarla distinto.",
        "Aun la hoja más cargada de cada rango queda cubierta por su precio (ver hoja Desglose, columna "
        "\"Costo máximo\").",
        "Precios de tóner y papel de septiembre de 2026. Si cambian, se actualizan en el programa (botón Costos) "
        "o en la hoja Supuestos de este archivo, y todo se recalcula solo.",
    ]
    for i, n in enumerate(notas, start=1):
        c = ws.cell(fila, 1, f"{i}. {n}")
        c.font = f(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
        ws.row_dimensions[fila].height = 14 * -(-len(n) // 110) + 4  # ~110 caracteres por renglón
        fila += 1

    for col, w in zip("ABCDEF", [17, 42, 13, 13, 15, 13]):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws["D4"].comment = Comment("Precios acordados en septiembre de 2026.", "Cotizador")


def main():
    wb = Workbook()
    wb.active.title = "Tabla de precios"
    s = hoja_supuestos(wb)
    filas, fila_toner = hoja_desglose(wb, s)
    hoja_tabla(wb, filas, fila_toner)
    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(SALIDA)


if __name__ == "__main__":
    main()
